"""
Minimal Markdown → Vairiot-themed Word converter.

Supports the subset used in the Vairiot Asset Intelligence documentation set: headings
(#–####), paragraphs, bullet and numbered lists (two levels), pipe tables,
fenced code blocks, block quotes (rendered as call-outs), and inline
**bold**, *italic*, `code` and [links](url).

Diagram handling: ```mermaid blocks are kept in the Markdown for GitHub and
skipped in Word. A comment of the form
    <!-- word-image: file.png | Caption text -->
inserts a rendered diagram image (path relative to the images directory).
"""

from __future__ import annotations

import re
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Mm, Pt

import vairiot_theme as vt

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)")
IMAGE_TAG = re.compile(r"<!--\s*word-image:\s*([^|]+?)\s*(?:\|\s*(.+?))?\s*-->")
NUMBERED = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
BULLET = re.compile(r"^(\s*)[-*]\s+(.*)$")


def add_inline(paragraph, text: str, size: float | None = None,
               colour: str | None = None, bold: bool = False) -> None:
    """Add text with inline Markdown formatting to a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            vt.set_run_font(r, "body", size, colour, True)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            vt.set_run_font(r, "mono", (size or 10) - 1, vt.VIOLET)
        elif part.startswith("[") and "](" in part:
            label = part[1:part.index("](")]
            r = paragraph.add_run(label)
            vt.set_run_font(r, "body", size, vt.MAUVE, bold)
            r.underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            vt.set_run_font(r, "body", size, colour, bold)
            r.italic = True
        else:
            r = paragraph.add_run(part)
            vt.set_run_font(r, "body", size, colour, bold or None)


def _split_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    cols = len(header)
    table = doc.add_table(rows=1, cols=cols)
    size = 8.5 if cols >= 4 else 9
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        vt.shade(cell, vt.VIOLET)
        vt.cell_margins(cell)
        add_inline(cell.paragraphs[0], text, size, "FFFFFF", True)
    vt.repeat_header(table.rows[0])
    for n, row in enumerate(body):
        cells = table.add_row().cells
        for i in range(cols):
            text = row[i] if i < len(row) else ""
            vt.cell_margins(cells[i])
            if n % 2 == 1:
                vt.shade(cells[i], vt.WASH)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, size)
    vt.table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code(doc, lines: list[str]) -> None:
    for i, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(4 if i == 0 else 0)
        pf.left_indent = Mm(3)
        vt.paragraph_shade(p, vt.WASH)
        r = p.add_run(line if line else " ")
        vt.set_run_font(r, "mono", 8, vt.CHARCOAL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, path: Path, caption: str | None) -> None:
    if not path.exists():
        p = doc.add_paragraph()
        add_inline(p, f"*Diagram not found: {path.name}*")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Mm(172))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        vt.set_run_font(r, "body", 8.5, vt.GREY)
        r.italic = True


def _list_item(doc, text: str, level: int, number: str | None) -> None:
    """Bullet or numbered item. Numbers are written explicitly so each list
    restarts at its own first number (Word's shared List Number counter would
    otherwise continue across the whole document)."""
    if number is None:
        p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
        add_inline(p, text)
        return
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Mm(7 + 6 * level)
    pf.first_line_indent = Mm(-7)
    pf.space_after = Pt(3)
    pf.tab_stops.add_tab_stop(Mm(7 + 6 * level))
    r = p.add_run(f"{number}.\t")
    vt.set_run_font(r, "head", 10, vt.MAUVE, True)
    add_inline(p, text)


def convert(doc, markdown: str, images_dir: Path,
            h1_page_break: bool = True, skip_first_h1: bool = False) -> None:
    """Render a Markdown string into the given document."""
    lines = markdown.splitlines()
    i, first_h1 = 0, True
    para: list[str] = []

    def flush() -> None:
        if para:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        img = IMAGE_TAG.search(stripped)
        if img:
            flush()
            add_image(doc, images_dir / img.group(1).strip(), img.group(2))
            i += 1
            continue
        if stripped.startswith("<!--"):
            i += 1
            continue

        if stripped.startswith("```"):
            flush()
            lang = stripped[3:].strip()
            block, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if lang != "mermaid":
                add_code(doc, block)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                if skip_first_h1 and first_h1:
                    first_h1 = False
                    i += 1
                    continue
                if h1_page_break and not first_h1:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first_h1 = False
            p = doc.add_paragraph(style=f"Heading {level}")
            p.add_run(re.sub(r"[`*]", "", text))
            if level == 1:
                vt.paragraph_rule(p)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) \
                and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            flush()
            rows = [_split_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            text = " ".join(quote)
            label, _, body = text.partition(":")
            label = re.sub(r"[*]", "", label) if len(label) < 24 else "Note"
            vt.callout(doc, label, body.strip() if len(label) < 24 and body else text)
            continue

        num, bul = NUMBERED.match(line), BULLET.match(line)
        if num or bul:
            flush()
            m = num or bul
            level = 1 if len(m.group(1)) >= 2 else 0
            text = m.group(3) if num else m.group(2)
            _list_item(doc, text, level, m.group(2) if num else None)
            i += 1
            continue

        if not stripped or stripped == "---":
            flush()
            i += 1
            continue

        para.append(line)
        i += 1
    flush()
