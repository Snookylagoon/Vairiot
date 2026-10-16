"""
Vairiot document theme for python-docx.

Applies the Vairiot design system (Montserrat headings and body, IBM Plex Mono
for codes; pink → mauve → violet gradient palette) to a Word document, and
provides the cover page, header/footer and table helpers used by every
Vairiot Asset Intelligence stage document.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BRAND_DIR = Path("/Volumes/DRSssd/Projects/GitHub/brand-assets/logos/vairiot")

PINK = "FF0DCC"
MAUVE = "A05B97"
VIOLET = "615AA0"
CHARCOAL = "2B3132"
WASH = "F8F0FA"
GREY = "6B7280"
RULE = "E5D9EA"

FONT_HEAD = "Montserrat ExtraBold"
FONT_BODY = "Montserrat"
FONT_MONO = "IBM Plex Mono"


# Schema-order successors, so raw elements are inserted where Word expects them
_TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                   "w:vAlign", "w:hideMark")
_TCPR_AFTER_MAR = ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
_TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                        "w:tblCaption", "w:tblDescription")
_PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                  "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                  "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection",
                  "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                  "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_PPR_AFTER_BDR = ("w:shd",) + _PPR_AFTER_SHD
_SETTINGS_AFTER_UPDATE = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr",
                          "w:compat", "w:docVars", "w:rsids",
                          "w:attachedSchema", "w:themeFontLang",
                          "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
                          "w:doNotAutoCompressPictures", "w:forceUpgrade",
                          "w:captions", "w:readModeInkLockDown",
                          "w:smartTagType", "w:shapeDefaults",
                          "w:doNotEmbedSmartTags", "w:decimalSymbol",
                          "w:listSeparator")


def rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _set_font(style_or_run, name: str, size: float | None = None,
              colour: str | None = None, bold: bool | None = None) -> None:
    font = style_or_run.font
    font.name = name
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        font.size = Pt(size)
    if colour is not None:
        font.color.rgb = rgb(colour)
    if bold is not None:
        font.bold = bold


def shade(cell, hex_fill: str) -> None:
    """Fill a table cell with a solid colour."""
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.insert_element_before(shd, *_TCPR_AFTER_SHD)


def cell_margins(cell, top=60, bottom=60, left=90, right=90) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom),
                      ("start", left), ("end", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.insert_element_before(mar, *_TCPR_AFTER_MAR)


def table_borders(table, colour: str = RULE, size: int = 4) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tblpr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def paragraph_rule(paragraph, colour: str = MAUVE, size: int = 12) -> None:
    """Bottom border under a paragraph (used beneath H1 headings)."""
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), colour)
    bdr.append(bottom)
    ppr.insert_element_before(bdr, *_PPR_AFTER_BDR)


def paragraph_shade(paragraph, hex_fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.insert_element_before(shd, *_PPR_AFTER_SHD)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    _set_font(normal, FONT_BODY, 10, CHARCOAL)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    heads = {"Title": (28, VIOLET), "Heading 1": (18, VIOLET),
             "Heading 2": (14, MAUVE), "Heading 3": (11.5, CHARCOAL),
             "Heading 4": (10.5, MAUVE)}
    for name, (size, colour) in heads.items():
        st = styles[name]
        _set_font(st, FONT_HEAD, size, colour, bold=True)
        st.font.italic = False
        pf = st.paragraph_format
        pf.space_before = Pt(16 if name == "Heading 1" else 10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        st = styles[name]
        _set_font(st, FONT_BODY, 10, CHARCOAL)
        st.paragraph_format.space_after = Pt(3)


def _page_setup(doc: Document) -> None:
    for section in doc.sections:
        section.page_height = Mm(297)
        section.page_width = Mm(210)
        section.orientation = WD_ORIENT.PORTRAIT
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(18)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
        section.header_distance = Mm(9)
        section.footer_distance = Mm(9)


def _field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _header_footer(doc: Document, doc_title: str) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    hp = section.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("VAIRIOT  ")
    _set_font(r, FONT_HEAD, 8, VIOLET, bold=True)
    r = hp.add_run("Asset Intelligence  ·  " + doc_title)
    _set_font(r, FONT_BODY, 8, GREY)
    paragraph_rule(hp, PINK, 6)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("Confidential  ·  VAIRIOT FZCO  ·  Page ")
    _set_font(r, FONT_BODY, 8, GREY)
    r = fp.add_run()
    _set_font(r, FONT_BODY, 8, GREY)
    _field(r, "PAGE")
    r = fp.add_run(" of ")
    _set_font(r, FONT_BODY, 8, GREY)
    r = fp.add_run()
    _set_font(r, FONT_BODY, 8, GREY)
    _field(r, "NUMPAGES")


def new_document(doc_title: str) -> Document:
    doc = Document()
    _page_setup(doc)
    _configure_styles(doc)
    _header_footer(doc, doc_title)
    return doc


def cover_page(doc: Document, title: str, subtitle: str,
               meta: list[tuple[str, str]]) -> None:
    """Gradient-banded cover page with the Vairiot logo and a metadata block."""
    logo = BRAND_DIR / "vairiot-logo-white.png"
    divider = BRAND_DIR / "vairiot-divider.png"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Mm(28)
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Mm(70))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Mm(30)
    r = p.add_run("ASSET INTELLIGENCE")
    _set_font(r, FONT_HEAD, 11, PINK, bold=True)
    r.font.all_caps = True

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    _set_font(r, FONT_HEAD, 30, VIOLET, bold=True)

    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    _set_font(r, FONT_BODY, 13, CHARCOAL)

    if divider.exists():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.add_run().add_picture(str(divider), width=Mm(174))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Mm(40)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in meta:
        row = table.add_row().cells
        for c in row:
            cell_margins(c)
        shade(row[0], WASH)
        rl = row[0].paragraphs[0].add_run(label.upper())
        _set_font(rl, FONT_HEAD, 8, MAUVE, bold=True)
        rv = row[1].paragraphs[0].add_run(value)
        _set_font(rv, FONT_BODY, 10, CHARCOAL)
    table.columns[0].width = Mm(45)
    for row in table.rows:
        row.cells[0].width = Mm(45)
        row.cells[1].width = Mm(129)
    table_borders(table, RULE, 4)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Mm(14)
    r = p.add_run("Prepared " + date.today().strftime("%d %B %Y")
                  + "  ·  British English  ·  Confidential")
    _set_font(r, FONT_BODY, 8, GREY)
    p.add_run().add_break(WD_BREAK.PAGE)


def contents_page(doc: Document) -> None:
    """Insert a Word table of contents field (updates on open)."""
    doc.add_paragraph("Contents", style="Heading 1")
    p = doc.add_paragraph()
    r = p.add_run()
    _field(r, 'TOC \\o "1-2" \\h \\z \\u')
    note = doc.add_paragraph()
    rn = note.add_run("Right-click the contents and choose Update Field "
                      "if page numbers are not shown.")
    _set_font(rn, FONT_BODY, 8, GREY)
    rn.italic = True
    note.add_run().add_break(WD_BREAK.PAGE)
    _enable_update_fields(doc)


def _enable_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.insert_element_before(el, *_SETTINGS_AFTER_UPDATE)


def callout(doc: Document, label: str, text: str,
            colour: str = PINK) -> None:
    """Single-cell shaded call-out box with a coloured left label."""
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade(cell, WASH)
    cell_margins(cell, 120, 120, 160, 160)
    p = cell.paragraphs[0]
    r = p.add_run(label.upper() + "   ")
    _set_font(r, FONT_HEAD, 8.5, colour, bold=True)
    r = p.add_run(text)
    _set_font(r, FONT_BODY, 9.5, CHARCOAL)
    table_borders(table, colour, 6)
    doc.add_paragraph()


def set_run_font(run, kind: str = "body", size: float | None = None,
                 colour: str | None = None, bold: bool | None = None) -> None:
    name = {"body": FONT_BODY, "head": FONT_HEAD, "mono": FONT_MONO}[kind]
    _set_font(run, name, size, colour, bold)
