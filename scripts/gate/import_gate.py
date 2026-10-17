"""W6 close on the gate stack (stage 2.14): the staged import through nginx,
worked by the worker image. 10,000 rows from upload to commit under 2 minutes
(test plan §3.6), error rows left out, update mode changes no asset number.

Usage: python3 scripts/gate/import_gate.py <gate dir> [rows] [finance-ref prefix]   (after phase1_gate.py)
"""
import io, json, socket, ssl, subprocess, sys, time, urllib.error, urllib.request, uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

G = Path(sys.argv[1])
ROWS = int(sys.argv[2]) if len(sys.argv) > 2 else 10_000
P = sys.argv[3] if len(sys.argv) > 3 else 'GF'
HOST, PORT = 'assets.gate.localhost', 18443
BASE = f'https://{HOST}:{PORT}/api/v1'
COMPOSE = ['docker', 'compose', '-p', 'vairiot-gate', '--env-file', str(G / 'gate.env'),
           '-f', 'docker-compose.prod.yml', '-f', 'docker-compose.standalone.yml', '-f', str(G / 'gate.override.yml')]
INFRA = str(Path(__file__).resolve().parents[2] / 'infra')
_orig = socket.getaddrinfo
socket.getaddrinfo = lambda h, *a, **k: _orig('127.0.0.1' if h.endswith('.localhost') else h, *a, **k)
CTX = ssl.create_default_context(); CTX.check_hostname = False; CTX.verify_mode = ssl.CERT_NONE

results = []
def gate(name, ok, detail=''):
    results.append({'gate': name, 'pass': bool(ok), 'detail': detail})
    print(('PASS ' if ok else 'FAIL ') + name + (f' — {detail}' if detail else ''), flush=True)

token = None
def call(method, path, body=None, raw=None, ctype='application/json'):
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    req = urllib.request.Request(BASE + path, data=data, method=method)
    if data is not None: req.add_header('Content-Type', ctype)
    if token: req.add_header('Authorization', f'Bearer {token}')
    try:
        with urllib.request.urlopen(req, context=CTX, timeout=120) as r:
            b = r.read(); return r.status, (json.loads(b) if r.headers.get_content_type() == 'application/json' else b)
    except urllib.error.HTTPError as e:
        b = e.read()
        try: return e.code, json.loads(b)
        except Exception: return e.code, b

def sql(q):
    return subprocess.run(COMPOSE + ['exec', '-T', 'postgres', 'psql', '-U', 'vairiot', '-d', 'vairiot', '-Atc', q],
                          cwd=INFRA, capture_output=True, text=True, check=True).stdout.strip()

def upload(name, rows, mode='CREATE'):
    wb = Workbook(); ws = wb.active; ws.title = 'Register'
    ws.append(['TUDA register — gate']);
    for r in rows: ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    b = uuid.uuid4().hex
    body = (f'--{b}\r\nContent-Disposition: form-data; name="mode"\r\n\r\n{mode}\r\n'
            f'--{b}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\n'
            'Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n').encode() + buf.getvalue() + f'\r\n--{b}--\r\n'.encode()
    return call('POST', '/imports', raw=body, ctype=f'multipart/form-data; boundary={b}')

def wait(batch, until, limit=600):
    t0 = time.time()
    while time.time() - t0 < limit:
        s, b = call('GET', f'/imports/{batch}')
        if s == 429: time.sleep(5); continue
        if b['status'] in until or b['status'] == 'FAILED': return b
        time.sleep(2)
    raise TimeoutError(batch)

s, b = call('POST', '/auth/login', {'email': 'gate-admin@tuda.example', 'password': (G / 'admin.password').read_text(), 'tenantId': 'tuda'})
assert s == 200, (s, b)
token = b['accessToken']
before = int(sql("SELECT count(*) FROM assets WHERE \"tenantId\"='tuda'"))

HEAD = ['Inv No', 'Name', 'Class', 'Cost', 'Bought', 'Ref', 'Grade']
MAPPING = {'fields': {'legacyAssetNumber': 'Inv No', 'name': 'Name', 'classCode': 'Class', 'purchaseCost': 'Cost', 'purchaseDate': 'Bought', 'financeRef': 'Ref', 'conditionGrade': 'Grade'},
           'transforms': {'dateFormat': 'DMY', 'decimalSeparator': ','}}
rows = [[f'GATE-{i}', f'Gate sign {i}', 'SGN', '1 000,00', '10.08.2026', f'{P}-{i}', 4] for i in range(ROWS)]
rows[5] = ['GATE-5', 'Gate sign 5', 'XYZ', '1 000,00', 'soon', f'{P}-5', 4]      # an error row: left out
rows[6][5] = f'{P}-0'                                                              # a duplicate finance reference

t0 = time.time()
s, b = upload('gate-register.xlsx', [HEAD, *rows])
gate('Upload through nginx answers 202', s == 202, f'{s} {b.get("status") if isinstance(b, dict) else b}')
batch = b['id']
b = wait(batch, {'PARSED'}); t_parse = time.time()
gate('The worker reads the file from MinIO', b['status'] == 'PARSED' and b['rowCount'] == ROWS and b['headerRow'] == 2, f"{b['status']}, {b['rowCount']} rows, header row {b['headerRow']}, {b.get('jobError')}")
s, _ = call('PUT', f'/imports/{batch}/mapping', {'mapping': MAPPING})
b = wait(batch, {'VALIDATED'}); t_valid = time.time()
gate('The worker checks every row', b['status'] == 'VALIDATED' and b['errorCount'] == 2 and b['validCount'] == ROWS - 2, f"valid {b['validCount']}, warnings {b['warningCount']}, errors {b['errorCount']}")
s, _ = call('POST', f'/imports/{batch}/commit')
b = wait(batch, {'COMMITTED'}); t_done = time.time()
total = t_done - t0
gate(f'{ROWS} rows from upload to commit under 2 minutes, through the worker and the internal API', b['status'] == 'COMMITTED' and b['committedCount'] == ROWS - 2 and total < 120,
     f"{b['status']} {b['committedCount']} committed; parse {t_parse - t0:.1f} s, validate {t_valid - t_parse:.1f} s, commit {t_done - t_valid:.1f} s, total {total:.1f} s; {b.get('jobError')}")
after = int(sql("SELECT count(*) FROM assets WHERE \"tenantId\"='tuda'"))
gate('Error rows are left out', after - before == ROWS - 2, f'{after - before} new assets')
refs = int(sql("SELECT count(*) FROM asset_secondary_identifiers WHERE \"tenantId\"='tuda' AND scheme='FINANCE_REF' AND value LIKE '" + P + "-%'"))
migrated = int(sql("SELECT count(*) FROM asset_events WHERE \"tenantId\"='tuda' AND \"eventType\"='MIGRATED'"))
gate('Each asset has its finance reference and a MIGRATED event', refs == ROWS - 2 and migrated >= ROWS - 2, f'{refs} references, {migrated} MIGRATED')
s, x = call('GET', f'/imports/{batch}/errors.xlsx')
ws = load_workbook(io.BytesIO(x)).active
gate('The error rows download as a workbook', s == 200 and ws.max_row == 3, f'{ws.max_row - 1} rows: {ws.cell(2, ws.max_column).value}')

# Update mode: ten rows by finance reference, new costs; asset numbers unchanged.
numbers = sql("SELECT string_agg(a.\"assetNumber\", ',' ORDER BY i.value) FROM assets a JOIN asset_secondary_identifiers i ON i.\"assetId\" = a.id WHERE i.value IN (" + ','.join(f"'{P}-{i}'" for i in (1, 2, 3, 4, 7, 8, 9, 10, 11, 12)) + ")")
s, b = upload('gate-update.xlsx', [['Ref', 'Cost'], *[[f'{P}-{i}', '1 250,00'] for i in (1, 2, 3, 4, 7, 8, 9, 10, 11, 12)]], 'UPDATE')
u = b['id']; wait(u, {'PARSED'})
call('PUT', f'/imports/{u}/mapping', {'mapping': {'fields': {'financeRef': 'Ref', 'purchaseCost': 'Cost'}, 'transforms': {'decimalSeparator': ','}}})
wait(u, {'VALIDATED'}); call('POST', f'/imports/{u}/commit'); b = wait(u, {'COMMITTED'})
numbers_after = sql("SELECT string_agg(a.\"assetNumber\", ',' ORDER BY i.value) FROM assets a JOIN asset_secondary_identifiers i ON i.\"assetId\" = a.id WHERE i.value IN (" + ','.join(f"'{P}-{i}'" for i in (1, 2, 3, 4, 7, 8, 9, 10, 11, 12)) + ")")
cost = sql("SELECT DISTINCT a.\"purchaseCost\" FROM assets a JOIN asset_secondary_identifiers i ON i.\"assetId\" = a.id WHERE i.value IN ('" + P + "-1','" + P + "-12')")
gate('Update mode by finance reference changes no asset number', b['committedCount'] == 10 and numbers == numbers_after and cost == '1250.00', f'{b["committedCount"]} updated, cost {cost}, numbers unchanged: {numbers == numbers_after}')

(G / 'import-results.json').write_text(json.dumps(results, indent=2))
print(f"{sum(r['pass'] for r in results)}/{len(results)} passed")
