#!/usr/bin/env python3
"""Generate scripts/samples/register-sample.xlsx: a synthetic legacy asset
register in the shape public-sector registers usually arrive in, with known
problems planted so profile-register.py has something to find.

Everything here is invented: no real assets, people or values.

    python3 scripts/samples/make-register-sample.py

Shape: a title row and a blank row above the real header (as exported
registers often have), bilingual Georgian/English headers, mixed number and
date formats. Planted problems (the expected findings are in PLANTED below):
  - duplicate inventory numbers
  - blank asset names
  - costs that aren't numbers ("n/a", "TBC", "12,5,0")
  - acquisition dates before 1990 or in the future
"""
from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).with_name("register-sample.xlsx")

HEADERS = [
    "ინვ. № / Inventory No",
    "დასახელება / Name",
    "კატეგორია / Category",
    "ქარხნული № / Serial No",
    "ლოკაცია / Location",
    "შეძენის თარიღი / Acquisition date",
    "ღირებულება (₾) / Cost (GEL)",
    "მდგომარეობა / Condition",
    "პასუხისმგებელი / Custodian",
]

CATEGORIES = {
    "Bus (12 m)": ("Bus", 380_000, 520_000),
    "Traffic signal controller": ("Signal controller", 9_000, 16_000),
    "Bus shelter": ("Shelter", 6_000, 11_000),
    "CCTV camera": ("Camera", 900, 2_400),
    "Desktop computer": ("PC", 1_400, 3_200),
    "Ticket validator": ("Validator", 1_100, 1_900),
}
LOCATIONS = ["Depot 1 — Didube", "Depot 2 — Gldani", "Head office", "Rustaveli Ave", "Vake", "Saburtalo"]
CONDITIONS = ["კარგი / good", "საშუალო / fair", "ცუდი / poor"]
CUSTODIANS = ["Fleet unit", "Traffic management", "IT department", "Stops & shelters"]

# Row numbers below are spreadsheet rows (header is row 3, data starts row 4).
PLANTED = {
    "duplicate_inventory_numbers": {"TUDA-000017": [20, 41], "TUDA-000052": [55, 61]},
    "blank_names": [12, 33, 47],
    "non_numeric_costs": {18: "n/a", 29: "TBC", 64: "12,5,0"},
    "dates_out_of_range": {26: "1987-06-30", 38: "2031-01-15"},
}


def main() -> None:
    rng = random.Random(20261006)  # fixed seed: the file is reproducible
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    ws["A1"] = "TUDA fixed asset register — export 2026 (SYNTHETIC SAMPLE)"
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(HEADERS)

    rows = 70
    for i in range(1, rows + 1):
        sheet_row = i + 3
        cat = rng.choice(list(CATEGORIES))
        short, lo, hi = CATEGORIES[cat]
        cost: object = round(rng.uniform(lo, hi), 2)
        # Legacy registers mix number formats: some costs typed as text.
        if i % 9 == 0:
            cost = f"{cost:,.2f}".replace(",", " ").replace(".", ",")   # "12 345,67"
        acquired: object = dt.date(2004, 1, 1) + dt.timedelta(days=rng.randint(0, 7600))
        if i % 7 == 0:
            acquired = acquired.strftime("%d.%m.%Y")                    # text date
        row = [
            f"TUDA-{i:06d}",
            f"{cat} #{i}",
            short,
            f"SN-{rng.randint(10_000_000, 99_999_999)}",
            rng.choice(LOCATIONS),
            acquired,
            cost,
            rng.choice(CONDITIONS),
            rng.choice(CUSTODIANS),
        ]
        ws.append(row)

        if sheet_row in PLANTED["blank_names"]:
            ws.cell(sheet_row, 2).value = None
        if sheet_row in PLANTED["non_numeric_costs"]:
            ws.cell(sheet_row, 7).value = PLANTED["non_numeric_costs"][sheet_row]
        if sheet_row in PLANTED["dates_out_of_range"]:
            ws.cell(sheet_row, 6).value = dt.date.fromisoformat(PLANTED["dates_out_of_range"][sheet_row])

    for number, sheet_rows in PLANTED["duplicate_inventory_numbers"].items():
        for r in sheet_rows:
            ws.cell(r, 1).value = number

    for col, width in zip("ABCDEFGHI", (16, 34, 18, 16, 20, 18, 16, 16, 20)):
        ws.column_dimensions[col].width = width
    wb.save(OUT)
    print(f"wrote {OUT} ({rows} assets)")


if __name__ == "__main__":
    main()
