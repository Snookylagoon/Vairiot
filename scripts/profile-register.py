#!/usr/bin/env python3
"""Profile a legacy asset register before importing it into Vairiot.

    python3 scripts/profile-register.py REGISTER.xlsx [--sheet NAME] [--header-row N]
                                        [--out REPORT.xlsx] [--mapping MAPPING.json]

Reads an Excel (.xlsx/.xlsm) or CSV register and writes:
  * an Excel profile report (A4 landscape): a summary, one row per column (fill
    rate, distinct values, samples, detected type), every data-quality issue
    with its spreadsheet row, and the suggested mapping;
  * a JSON column mapping for the S2 importer (register column -> Vairiot
    asset field, with a confidence), for a person to confirm before import.

Data-quality checks (on the columns the mapping identifies):
  * duplicate asset/inventory numbers
  * blank asset names
  * costs that are not numbers (European "1 234,50", "₾100" and "1,234.50" are
    numbers; "n/a", "TBC" are not)
  * dates outside 1990-01-01 … today (Excel dates, ISO, dd.mm.yyyy, dd/mm/yyyy)

Headers in English, Georgian and Russian are recognised. Values are read as
text (so "00123" keeps its zeros) and typed by this script, not by pandas.
Requires: pandas, openpyxl (scripts/requirements.txt).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

DATE_MIN = dt.date(1990, 1, 1)
SAMPLE_VALUES = 5

# ─── Vairiot asset fields and the header words that suggest them ───────────────
# Keys are the asset importer's field names (vairiot-api import.service.ts
# ImportRow; the web import page maps CSV headers to the same keys). Synonyms
# are matched against normalised headers (lower case, no punctuation); a
# bilingual header such as "ინვ. № / Inventory No" is matched on each part.
FIELD_SYNONYMS: dict[str, list[str]] = {
    "name": [
        "name", "asset name", "description of asset", "item", "title",
        "დასახელება", "სახელწოდება", "наименование", "название",
    ],
    "description": ["description", "details", "აღწერა", "описание"],
    "categoryName": ["category", "class", "asset class", "type", "group", "კატეგორია", "ტიპი", "категория", "группа", "тип"],
    "siteName": ["site", "facility", "building", "depot", "ობიექტი", "объект", "площадка"],
    "serialNumber": [
        "serial number", "serial no", "serial", "s n", "factory number",
        "ქარხნული ნომერი", "ქარხნული", "სერიული ნომერი", "заводской номер", "серийный номер",
    ],
    "modelNumber": ["model", "model number", "მოდელი", "модель"],
    "manufacturer": ["manufacturer", "make", "brand", "მწარმოებელი", "производитель"],
    "barcode": ["barcode", "ean", "manufacturer ean", "შტრიხკოდი", "штрихкод"],
    "rfidTag": ["rfid", "rfid tag", "epc"],
    "purchaseCost": [
        "cost", "purchase cost", "acquisition cost", "historical cost", "value", "price", "book value",
        "ღირებულება", "საწყისი ღირებულება", "стоимость", "первоначальная стоимость", "цена",
    ],
    "purchaseDate": [
        "purchase date", "acquisition date", "date acquired", "date of purchase", "acquired",
        "შეძენის თარიღი", "შეძენა", "дата приобретения", "дата покупки", "дата ввода",
    ],
    "supplier": ["supplier", "vendor", "მომწოდებელი", "поставщик"],
    "warrantyExpiry": ["warranty expiry", "warranty", "warranty end", "საგარანტიო", "гарантия"],
    "condition": ["condition", "state", "მდგომარეობა", "состояние"],
    "status": ["status", "სტატუსი", "статус"],
    "notes": ["notes", "remarks", "comment", "შენიშვნა", "примечание"],
    "purchaseOrderNumber": ["po number", "purchase order", "order number", "შეკვეთის ნომერი", "номер заказа"],
    "invoiceNumber": ["invoice number", "invoice", "ინვოისი", "ზედნადები", "номер счета", "накладная"],
    "residualValue": ["residual value", "salvage value", "ნარჩენი ღირებულება", "остаточная стоимость"],
    "usefulLifeMonths": ["useful life", "useful life months", "life months", "სასარგებლო ვადა", "срок полезного использования"],
    "depreciationMethod": ["depreciation method", "ცვეთის მეთოდი", "метод амортизации"],
    "depreciationStartDate": ["depreciation start", "depreciation start date", "in service date", "ექსპლუატაციაში შესვლა"],
    # Recognised, but the importer has no field for them yet: reported, not
    # silently dropped.
    "legacyAssetNumber": [
        "asset number", "asset no", "inventory number", "inventory no", "inv no", "asset id",
        "asset tag", "tag number", "register number",
        "ინვენტარის ნომერი", "ინვ", "ინვენტარული ნომერი", "საინვენტარო ნომერი",
        "инвентарный номер", "инв номер", "инв",
    ],
    "locationName": ["location", "room", "place", "ლოკაცია", "ადგილმდებარეობა", "местонахождение", "расположение"],
    "custodian": ["custodian", "responsible", "responsible person", "owner", "პასუხისმგებელი", "ответственный", "мол"],
}

NOT_IMPORTABLE = {
    "legacyAssetNumber": "Vairiot allocates its own asset numbers; keep this one (S2 import will carry it as the legacy reference). Until then map it to notes if needed.",
    "locationName": "The importer sets the site only; locations are added after import (or in S2).",
    "custodian": "No custodian field yet; map to notes if needed.",
}

# What each check needs from the mapping.
NUMBER_FIELDS = {"purchaseCost", "residualValue"}
DATE_FIELDS = {"purchaseDate", "depreciationStartDate"}   # warranty expiry may rightly be in the future


# ─── Reading ───────────────────────────────────────────────────────────────────
def read_raw(path: Path, sheet: str | None) -> tuple[pd.DataFrame, str, str]:
    """All cells as Python objects (no header), plus the sheet used and how it was read."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        with pd.ExcelFile(path, engine="openpyxl") as book:
            name = sheet or book.sheet_names[0]
            if name not in book.sheet_names:
                raise SystemExit(f"sheet {name!r} not found; sheets: {', '.join(book.sheet_names)}")
            # keep_default_na=False: pandas would otherwise turn "n/a", "NA", "null",
            # "None"… into blanks, hiding exactly the values this tool must report.
            frame = pd.read_excel(book, sheet_name=name, header=None, dtype=object, keep_default_na=False)
        return frame, name, "Excel"
    if suffix in {".csv", ".txt"}:
        for encoding in ("utf-8-sig", "utf-16", "cp1251", "latin-1"):
            try:
                frame = pd.read_csv(path, header=None, dtype=str, sep=None, engine="python",
                                    encoding=encoding, keep_default_na=False)
                return frame, path.name, f"CSV ({encoding})"
            except (UnicodeError, pd.errors.ParserError):
                continue
        raise SystemExit("could not read the CSV in UTF-8, UTF-16, Windows-1251 or Latin-1")
    raise SystemExit(f"unsupported file type {suffix!r}: use .xlsx, .xlsm or .csv")


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    return isinstance(value, str) and value.strip() == ""


def find_header_row(raw: pd.DataFrame, scan: int = 20) -> int:
    """Index of the header: the first row, among the top `scan`, that fills the
    most columns with text. Exported registers often put a title above it."""
    best, best_count = 0, -1
    for idx in range(min(scan, len(raw))):
        cells = raw.iloc[idx].tolist()
        text = sum(1 for v in cells if isinstance(v, str) and v.strip())
        if text > best_count:
            best, best_count = idx, text
    return best


# ─── Parsing values ────────────────────────────────────────────────────────────
NUMBER_JUNK = re.compile(r"[₾$€£\s ]|gel|lari|лари|ლარი", re.IGNORECASE)


def parse_number(value: Any) -> float | None:
    """A number from a cell, accepting "1 234,50", "1,234.50", "1.234,50" and "₾100"."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and not pd.isna(value):
        return float(value)
    if not isinstance(value, str):
        return None
    text = NUMBER_JUNK.sub("", value.strip())
    if not text or not re.fullmatch(r"-?[\d.,]+", text):
        return None
    if "," in text and "." in text:
        # The last separator is the decimal point.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) == 2 and len(parts[1]) != 3:
            text = text.replace(",", ".")       # decimal comma: "12,5"
        elif all(len(p) == 3 for p in parts[1:]):
            text = text.replace(",", "")        # thousands: "1,234,567"
        else:
            return None                         # "12,5,0" — ambiguous
    elif text.count(".") > 1:
        parts = text.split(".")
        if all(len(p) == 3 for p in parts[1:]):
            text = text.replace(".", "")        # "1.234.567"
        else:
            return None
    try:
        return float(text)
    except ValueError:
        return None


DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%y", "%Y.%m.%d")


def parse_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, (int, float)) and not isinstance(value, bool) and not pd.isna(value):
        # Excel serial dates (days since 1899-12-30) for 1900 … 2100.
        if 1 <= value <= 73050:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(value))
        return None
    if isinstance(value, str):
        text = value.strip().split(" ")[0]
        for fmt in DATE_FORMATS:
            try:
                return dt.datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def detect_type(values: list[Any]) -> tuple[str, float]:
    """The dominant type of the non-blank values and the share that fit it."""
    if not values:
        return "empty", 1.0
    n = len(values)
    dates = sum(1 for v in values if parse_date(v) is not None and not _plain_number(v))
    numbers = sum(1 for v in values if parse_number(v) is not None)
    ints = sum(1 for v in values if (num := parse_number(v)) is not None and float(num).is_integer())
    bools = sum(1 for v in values if isinstance(v, bool) or str(v).strip().lower() in {"yes", "no", "true", "false", "y", "n"})
    for name, count in (("date", dates), ("integer", ints if ints == numbers else 0), ("number", numbers), ("yes/no", bools)):
        if count / n >= 0.8:
            return name, count / n
    return ("text", 1.0) if max(dates, numbers, bools) / n < 0.2 else ("mixed", 1 - max(dates, numbers, bools) / n)


def _plain_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def show(value: Any, limit: int = 40) -> str:
    if isinstance(value, (dt.datetime, pd.Timestamp)):
        value = value.date()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value)
    return text if len(text) <= limit else text[: limit - 1] + "…"


# ─── Mapping ───────────────────────────────────────────────────────────────────
def normalise(text: str) -> str:
    text = unicodedata.normalize("NFKC", str(text)).lower()
    text = re.sub(r"\(.*?\)|№|#", " ", text)
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def score(header: str, synonym: str) -> float:
    if header == synonym:
        return 1.0
    if re.search(rf"(^| ){re.escape(synonym)}( |$)", header):
        return 0.9 if len(synonym) > 3 else 0.75
    return SequenceMatcher(None, header, synonym).ratio() * 0.85


def suggest_mapping(headers: list[str]) -> list[dict[str, Any]]:
    """Best Vairiot field for each header; each field is used at most once."""
    candidates = []
    for col, header in enumerate(headers):
        parts = [normalise(p) for p in re.split(r"\s/\s|\|", header)] + [normalise(header)]
        for fld, words in FIELD_SYNONYMS.items():
            best = max((score(p, w), w) for p in parts if p for w in words) if any(parts) else (0.0, "")
            if best[0] >= 0.7:
                candidates.append((best[0], col, fld, best[1]))
    taken_cols: set[int] = set()
    taken_fields: set[str] = set()
    chosen: dict[int, tuple[str, float, str]] = {}
    for conf, col, fld, word in sorted(candidates, key=lambda c: -c[0]):
        if col in taken_cols or fld in taken_fields:
            continue
        chosen[col] = (fld, conf, word)
        taken_cols.add(col)
        taken_fields.add(fld)
    result = []
    for col, header in enumerate(headers):
        if col in chosen:
            fld, conf, word = chosen[col]
            entry = {"column": header, "field": fld, "importable": fld not in NOT_IMPORTABLE,
                     "confidence": round(conf, 2), "reason": f'header matches "{word}"'}
            if fld in NOT_IMPORTABLE:
                entry["note"] = NOT_IMPORTABLE[fld]
            result.append(entry)
        else:
            result.append({"column": header, "field": None, "importable": False, "confidence": 0.0,
                           "reason": "no match — map by hand or skip"})
    return result


# ─── Profiling ─────────────────────────────────────────────────────────────────
@dataclass
class Issue:
    check: str
    row: int          # spreadsheet row number, as the user sees it
    column: str
    value: str
    detail: str


@dataclass
class Profile:
    source: str
    sheet: str
    read_as: str
    header_row: int
    rows: int
    columns: list[dict[str, Any]] = field(default_factory=list)
    mapping: list[dict[str, Any]] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)


def profile(path: Path, sheet: str | None = None, header_row: int | None = None,
            today: dt.date | None = None) -> Profile:
    today = today or dt.date.today()
    raw, sheet_name, read_as = read_raw(path, sheet)
    raw = raw.dropna(axis=1, how="all")
    h = (header_row - 1) if header_row else find_header_row(raw)
    headers = [show(v, 200) if not is_blank(v) else f"(column {get_column_letter(i + 1)})"
               for i, v in enumerate(raw.iloc[h].tolist())]
    body = raw.iloc[h + 1:]
    body = body[~body.apply(lambda r: all(is_blank(v) for v in r), axis=1)]
    # Spreadsheet row of each remaining data row (blank rows dropped above).
    sheet_rows = [int(i) + 1 for i in body.index]

    p = Profile(source=path.name, sheet=sheet_name, read_as=read_as, header_row=h + 1, rows=len(body))
    p.mapping = suggest_mapping(headers)
    field_of = {i: m["field"] for i, m in enumerate(p.mapping)}

    for i, header in enumerate(headers):
        values = body.iloc[:, i].tolist()
        filled = [v for v in values if not is_blank(v)]
        kind, fit = detect_type(filled)
        common = Counter(show(v) for v in filled).most_common(SAMPLE_VALUES)
        info: dict[str, Any] = {
            "column": header,
            "letter": get_column_letter(i + 1),
            "field": field_of[i],
            "filled": len(filled),
            "fill_rate": len(filled) / len(values) if values else 0.0,
            "distinct": len({show(v, 1000) for v in filled}),
            "type": kind,
            "type_fit": fit,
            "samples": ", ".join(f"{v} ({n})" if n > 1 else v for v, n in common),
            "max_length": max((len(show(v, 10_000)) for v in filled), default=0),
            "min": "", "max": "",
        }
        if kind in {"number", "integer"}:
            nums = [n for v in filled if (n := parse_number(v)) is not None]
            info["min"], info["max"] = show(min(nums)), show(max(nums))
        elif kind == "date":
            ds = [d for v in filled if (d := parse_date(v)) is not None]
            info["min"], info["max"] = min(ds).isoformat(), max(ds).isoformat()
        p.columns.append(info)

    def col_of(fld: str) -> int | None:
        return next((i for i, f in field_of.items() if f == fld), None)

    # Duplicate asset numbers.
    if (c := col_of("legacyAssetNumber")) is not None:
        seen: dict[str, list[int]] = {}
        for r, v in zip(sheet_rows, body.iloc[:, c].tolist()):
            if not is_blank(v):
                seen.setdefault(show(v, 1000).strip(), []).append(r)
        for value, rows in seen.items():
            if len(rows) > 1:
                for r in rows:
                    others = ", ".join(str(o) for o in rows if o != r)
                    p.issues.append(Issue("Duplicate asset number", r, headers[c], value, f"also on row(s) {others}"))
        for r, v in zip(sheet_rows, body.iloc[:, c].tolist()):
            if is_blank(v):
                p.issues.append(Issue("Blank asset number", r, headers[c], "", "every asset needs a number"))

    # Blank names.
    if (c := col_of("name")) is not None:
        for r, v in zip(sheet_rows, body.iloc[:, c].tolist()):
            if is_blank(v):
                p.issues.append(Issue("Blank asset name", r, headers[c], "", "name is required on import"))

    # Costs that aren't numbers.
    for fld in NUMBER_FIELDS:
        if (c := col_of(fld)) is not None:
            for r, v in zip(sheet_rows, body.iloc[:, c].tolist()):
                if not is_blank(v) and parse_number(v) is None:
                    p.issues.append(Issue("Cost is not a number", r, headers[c], show(v), "fix or clear before import"))
                elif (n := parse_number(v)) is not None and n < 0:
                    p.issues.append(Issue("Negative cost", r, headers[c], show(v), "costs can't be negative"))

    # Dates outside the plausible range.
    for fld in DATE_FIELDS:
        if (c := col_of(fld)) is not None:
            for r, v in zip(sheet_rows, body.iloc[:, c].tolist()):
                if is_blank(v):
                    continue
                d = parse_date(v)
                if d is None:
                    p.issues.append(Issue("Unreadable date", r, headers[c], show(v), "use YYYY-MM-DD or DD.MM.YYYY"))
                elif d < DATE_MIN or d > today:
                    p.issues.append(Issue("Date out of range", r, headers[c], d.isoformat(),
                                          f"outside {DATE_MIN.isoformat()} … {today.isoformat()}"))

    p.issues.sort(key=lambda i: (i.check, i.row))
    return p


# ─── Report ────────────────────────────────────────────────────────────────────
NAVY = "1F3A5F"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
ZEBRA_FILL = PatternFill("solid", fgColor="F3F6FA")
WARN_FILL = PatternFill("solid", fgColor="FDECEA")
OK_FILL = PatternFill("solid", fgColor="E8F5E9")
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _table(ws: Worksheet, start_row: int, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    for r, row in enumerate(rows, start_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if (r - start_row) % 2 == 0:
                cell.fill = ZEBRA_FILL
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    ws.print_title_rows = f"{start_row}:{start_row}"
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def _page_setup(ws: Worksheet, title: str) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddHeader.left.text = f"Register profile — {title}"
    ws.oddFooter.left.text = "&F"
    ws.oddFooter.right.text = "Page &P of &N"


def write_report(p: Profile, out: Path, mapping_file: Path) -> None:
    wb = Workbook()

    s = wb.active
    s.title = "Summary"
    s["A1"] = "Asset register profile"
    s["A1"].font = Font(bold=True, size=16, color=NAVY)
    counts = Counter(i.check for i in p.issues)
    mapped = sum(1 for m in p.mapping if m["field"])
    facts = [
        ("Source file", p.source), ("Sheet", p.sheet), ("Read as", p.read_as),
        ("Header row", p.header_row), ("Data rows", p.rows), ("Columns", len(p.columns)),
        ("Columns mapped to Vairiot fields", f"{mapped} of {len(p.mapping)}"),
        ("…of which importable today", sum(1 for m in p.mapping if m["importable"])),
        ("Profiled", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Mapping file", mapping_file.name),
    ]
    for r, (k, v) in enumerate(facts, 3):
        s.cell(r, 1, k).font = Font(bold=True)
        s.cell(r, 2, v)
    r0 = len(facts) + 4
    s.cell(r0, 1, "Data-quality findings").font = Font(bold=True, size=13, color=NAVY)
    checks = ["Duplicate asset number", "Blank asset number", "Blank asset name", "Cost is not a number",
              "Negative cost", "Unreadable date", "Date out of range"]
    rows = [[c, counts.get(c, 0)] for c in checks]
    _table(s, r0 + 1, ["Check", "Rows affected"], rows, [36, 60])
    s.freeze_panes = None
    for r in range(r0 + 2, r0 + 2 + len(rows)):
        s.cell(r, 2).fill = WARN_FILL if s.cell(r, 2).value else OK_FILL
    s.cell(r0 + len(rows) + 3, 1,
           "Rows are spreadsheet row numbers in the source file. Fix the source, re-run, then import.").font = Font(italic=True)
    _page_setup(s, p.source)

    c = wb.create_sheet("Columns")
    rows = [[col["letter"], col["column"], col["field"] or "—", f'{col["fill_rate"]:.0%}', col["filled"],
             col["distinct"], col["type"] + ("" if col["type_fit"] >= 0.999 else f' ({col["type_fit"]:.0%})'),
             col["min"], col["max"], col["max_length"], col["samples"]] for col in p.columns]
    _table(c, 1, ["Col", "Header", "Maps to", "Filled", "Values", "Distinct", "Detected type", "Min", "Max",
                  "Max length", "Most common values (count)"], rows, [6, 30, 15, 8, 8, 9, 14, 12, 12, 9, 60])
    for r in range(2, len(rows) + 2):
        c.cell(r, 4).alignment = Alignment(horizontal="right", vertical="top")
    _page_setup(c, p.source)

    i = wb.create_sheet("Issues")
    rows = [[x.check, x.row, x.column, x.value, x.detail] for x in p.issues] or [["No issues found", "", "", "", ""]]
    _table(i, 1, ["Check", "Row", "Column", "Value", "Detail"], rows, [24, 7, 34, 26, 40])
    _page_setup(i, p.source)

    m = wb.create_sheet("Mapping")
    rows = [[x["column"], x["field"] or "—",
             "yes" if x["importable"] else ("not yet" if x["field"] else "—"),
             x["confidence"] or "", x.get("note") or x["reason"]] for x in p.mapping]
    _table(m, 1, ["Register column", "Vairiot field", "Importable", "Confidence", "Why / note"], rows, [36, 20, 11, 11, 60])
    _page_setup(m, p.source)

    wb.save(out)


def write_mapping(p: Profile, path: Path) -> None:
    doc = {
        "source": p.source,
        "sheet": p.sheet,
        "headerRow": p.header_row,
        "dataRows": p.rows,
        "generatedAt": dt.datetime.now().isoformat(timespec="seconds"),
        "note": "Suggested by scripts/profile-register.py. Review every mapping before importing.",
        "mappings": p.mapping,
    }
    path.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("register", type=Path, help="register file (.xlsx, .xlsm or .csv)")
    ap.add_argument("--sheet", help="sheet to profile (default: the first)")
    ap.add_argument("--header-row", type=int, help="spreadsheet row of the column headers (default: detected)")
    ap.add_argument("--out", type=Path, help="report path (default: <register>-profile.xlsx)")
    ap.add_argument("--mapping", type=Path, help="mapping JSON path (default: <register>-mapping.json)")
    args = ap.parse_args(argv)
    if not args.register.exists():
        ap.error(f"{args.register} not found")

    out = args.out or args.register.with_name(f"{args.register.stem}-profile.xlsx")
    mapping = args.mapping or args.register.with_name(f"{args.register.stem}-mapping.json")
    p = profile(args.register, args.sheet, args.header_row)
    write_report(p, out, mapping)
    write_mapping(p, mapping)

    counts = Counter(i.check for i in p.issues)
    print(f"{p.source} [{p.sheet}]: {p.rows} rows, {len(p.columns)} columns, header on row {p.header_row}")
    print(f"  mapped {sum(1 for m in p.mapping if m['field'])} of {len(p.mapping)} columns")
    for check, n in sorted(counts.items()):
        print(f"  {check}: {n} row(s)")
    if not counts:
        print("  no data-quality issues found")
    print(f"report:  {out}\nmapping: {mapping}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
